"""
Excel 工具 - 精簡版（只負責匯出）
工號放最後一欄
"""
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.drawing.image import Image as XLImage
import os
from io import BytesIO


class ExcelHelper:

    @staticmethod
    def _make_header(ws):
        headers = ['#', '日期', 'Model', '站點', 'LINE', 'Sheet ID',
                   'T/C', 'X (cm)', 'Y (cm)', '程度 (LV)', 'Note', '照片', '工號']
        ws.append(headers)
        hfill  = PatternFill(start_color='1a6bb5', end_color='1a6bb5', fill_type='solid')
        hfont  = Font(bold=True, color='FFFFFF', size=11)
        halign = Alignment(horizontal='center', vertical='center')
        for cell in ws[1]:
            cell.fill  = hfill
            cell.font  = hfont
            cell.alignment = halign
        ws.row_dimensions[1].height = 20
        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 16
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 8
        ws.column_dimensions['H'].width = 10
        ws.column_dimensions['I'].width = 10
        ws.column_dimensions['J'].width = 12
        ws.column_dimensions['K'].width = 20
        ws.column_dimensions['L'].width = 18   # 照片
        ws.column_dimensions['M'].width = 12   # ★ 工號（最後欄）

    @staticmethod
    def export_to_bytes(records):
        """將紀錄清單匯出為 xlsx bytes（記憶體生成，不存磁碟）"""
        wb = Workbook()
        ws = wb.active
        ws.title = '背汙檢匯出'
        ExcelHelper._make_header(ws)

        for idx, rec in enumerate(records, start=1):
            row_num    = idx + 1
            photo_urls = rec.get('photoUrls', [])

            ws.append([
                idx,
                rec.get('date', ''),
                rec.get('model', ''),
                rec.get('station', ''),
                rec.get('line', ''),
                rec.get('sheetId', ''),
                rec.get('tc', ''),
                rec.get('x', ''),
                rec.get('y', ''),
                rec.get('lv', ''),
                rec.get('note', ''),
                '',                          # 照片欄（用圖片物件填）
                rec.get('empId', '')         # ★ 工號（最後欄）
            ])

            # 嵌入第一張照片（BytesIO 讀取，支援 UNC 路徑）
            img_ok = False
            for path in photo_urls:
                if not path or not os.path.exists(path):
                    continue
                try:
                    with open(path, 'rb') as f:
                        img_bytes = BytesIO(f.read())
                    img        = XLImage(img_bytes)
                    img.width  = 120
                    img.height = 90
                    ws.add_image(img, f'L{row_num}')
                    ws.row_dimensions[row_num].height = 70
                    img_ok = True
                    break
                except Exception as e:
                    print(f'⚠️ 圖片嵌入失敗 {path}：{e}')

            if not img_ok:
                ws.cell(row=row_num, column=12).value = '; '.join(photo_urls)
                ws.row_dimensions[row_num].height = 20

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
