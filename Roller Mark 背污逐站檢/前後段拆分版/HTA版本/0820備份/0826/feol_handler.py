#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
feol_handler.py — FEOL API 處理器
維護人員：FEOL Team
負責路由：/api/*（不含 /api/beol/*）

設計說明：
  每個函數的第一個參數為 handler（BeiwuAPIHandler 實例），
  由 APP.PY 路由器呼叫時傳入 self，例如：
      feol_handler.search_records(self, params)

【修改說明 v2】
  - download_template：改用網路路徑讀取範本，與 BEOL 一致
    路徑：\\tw100049089\ML8BC1\逐站檢資料\FEOL SAMPLE 01.xlsm
"""

import os
from datetime import datetime
from urllib.parse import unquote

from config_feol import TXT_DIR, PHOTOS_DIR, MACHINE_TXT
from utils.txt_helper     import TxtHelper
from utils.network_path   import decode_base64_image, save_photo_from_bytes
from utils.machine_helper import MachineHelper

# ===== 模組層級 DB 初始化 =====
txt_db     = TxtHelper(TXT_DIR)
machine_db = MachineHelper(MACHINE_TXT)

def log(msg, level='INFO'):
    print(f'[{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}] [FEOL/{level}] {msg}')

# ===== 工具函數 =====

def _normalize_station(st):
    if not st: return st
    STATIONS = ['PI 前', 'PI 後', 'PA 後', 'ODF 後', 'AGX 後', 'BS ITO 後']
    s = st.strip().upper().replace(' ', '')
    for canonical in STATIONS:
        if canonical.upper().replace(' ', '') == s:
            return canonical
    return st.strip()

STATION_ORDER = ['PI 前', 'PI 後', 'PA 後', 'ODF 後']

def _parse_y_range(y_val):
    import re
    s = str(y_val).strip() if y_val else ''
    if not s: return None, None
    m = re.match(r'^(\d+(?:\.\d+)?)[~\-](\d+(?:\.\d+)?)$', s)
    if m: return float(m.group(1)), float(m.group(2))
    try: v = float(s); return v, v
    except: return None, None

# ===== API 函數 =====

def health_check(handler):
    handler.send_json({'status': 'ok', 'timestamp': datetime.now().isoformat(),
        'paths': {'TXT_DIR': TXT_DIR, 'PHOTOS_DIR': PHOTOS_DIR, 'MACHINE_TXT': MACHINE_TXT,
                  'TXT_OK': os.path.exists(TXT_DIR), 'PHOTOS_OK': os.path.exists(PHOTOS_DIR),
                  'MACHINE_OK': machine_db.is_loaded()}})

def machine_status(handler):
    handler.send_json({'success': True, 'loaded': machine_db.is_loaded(),
                       'count': len(machine_db.records), 'headers': machine_db.headers})

def machine_match(handler, data):
    try:
        if not machine_db.is_loaded():
            handler.send_json({'success': False, 'error': '機台資料未載入'}, 503); return
        records   = data.get('records', [])
        tolerance = float(data.get('tolerance', 0.5))
        if not records and 'yValues' in data:
            records = [{'idx': i, 'y': y, 'line': '', 'station': ''} for i, y in enumerate(data['yValues'])]
        if not records:
            handler.send_json({'success': True, 'results': {}, 'headers': machine_db.headers}); return
        log(f'機台比對：{len(records)} 筆，tolerance={tolerance}cm')
        results = machine_db.batch_match(records, tolerance)
        handler.send_json({'success': True, 'results': results, 'headers': machine_db.headers})
    except Exception as e:
        log(f'❌ 機台比對失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def download_template(handler):
    # ★ 改用網路路徑，與 BEOL_SAMPLE_01.xlsm 相同做法
    template_path = r'\\tw100049089\ML8BC1\逐站檢資料\FEOL SAMPLE 01.xlsm'
    try:
        if not os.path.exists(template_path):
            handler.send_json({'success': False, 'error': '範例檔案不存在'}, 404); return
        with open(template_path, 'rb') as f:
            data = f.read()
        handler.send_response(200)
        handler.send_header('Content-Type', 'application/vnd.ms-excel.sheet.macroEnabled.12')
        handler.send_header('Content-Disposition', 'attachment; filename="FEOL SAMPLE 01.xlsm"')
        handler.send_header('Content-Length', str(len(data)))
        handler.send_header('Access-Control-Allow-Origin', '*')
        handler.end_headers()
        try:
            handler.wfile.write(data)
            log('✅ FEOL 範例 Excel 已下載')
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError, OSError):
            log('⚠️ 下載連線中斷')
    except Exception as e:
        log(f'❌ 範例下載失敗：{e}')

def search_records(handler, params):
    try:
        filters = {k: params[k] for k in ['model', 'station', 'line', 'tc', 'lv', 'empId'] if params.get(k)}
        records = txt_db.search(filters, params.get('dateFrom', ''), params.get('dateTo', ''))
        sid = params.get('sheetId', '').strip()
        if sid:
            records = [r for r in records if sid.lower() in r.get('sheetId', '').lower()]
        records.sort(key=lambda r: (r.get('date', ''), r.get('createdAt', '')), reverse=True)
        handler.send_json({'success': True, 'records': records})
    except Exception as e:
        log(f'❌ 搜尋失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def export_excel(handler, params):
    try:
        filters = {k: params[k] for k in ['model', 'station', 'line', 'tc', 'lv', 'empId'] if params.get(k)}
        records = txt_db.search(filters, params.get('dateFrom', ''), params.get('dateTo', ''))
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = '背汙檢'
        ws.append(['日期', 'Model', '站點', 'LINE', 'Sheet ID', 'T/C側', 'X(cm)', 'Y(cm)', '程度(LV)', 'Note', '工號'])
        for r in records:
            ws.append([r.get('date',''), r.get('model',''), r.get('station',''), r.get('line',''),
                       r.get('sheetId',''), r.get('tc',''), r.get('x',''), r.get('y',''),
                       r.get('lv',''), r.get('note',''), r.get('empId','')])
        from io import BytesIO
        buf = BytesIO(); wb.save(buf); buf.seek(0); data = buf.getvalue()
        handler.send_response(200)
        handler.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        handler.send_header('Content-Disposition', f'attachment; filename="背汙檢匯出_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"')
        handler.send_header('Content-Length', len(data))
        handler.send_header('Access-Control-Allow-Origin', '*')
        handler.end_headers()
        handler.wfile.write(data)
    except Exception as e:
        log(f'❌ 匯出失敗：{e}'); handler.send_error(500)

def serve_photo(handler, params):
    path = unquote(params.get('path', ''))
    try:
        if not path or not os.path.exists(path):
            handler.send_json({'error': '照片不存在'}, 404); return
        ext = path.rsplit('.', 1)[-1].lower()
        ct  = {'jpg': 'image/jpeg', 'jpeg': 'image/jpeg', 'png': 'image/png', 'gif': 'image/gif'}.get(ext, 'image/jpeg')
        with open(path, 'rb') as f: data = f.read()
        handler.send_response(200)
        handler.send_header('Content-Type', ct)
        handler.send_header('Content-Length', len(data))
        handler.send_header('Access-Control-Allow-Origin', '*')
        handler.end_headers()
        handler.wfile.write(data)
    except Exception as e:
        log(f'❌ 照片讀取失敗：{e}'); handler.send_error(500)

def save_records(handler, data):
    try:
        date    = data.get('date', '')
        records = data.get('records', [])
        if not date or not records:
            handler.send_json({'success': False, 'error': '日期或紀錄缺失'}, 400); return
        saved = 0
        for rec in records:
            photo_urls = []
            for photo in rec.get('photos', []):
                if 'dataUrl' not in photo: continue
                raw = decode_base64_image(photo['dataUrl'])
                if raw:
                    ok, path = save_photo_from_bytes(raw, photo['name'], PHOTOS_DIR, date)
                    if ok: photo_urls.append(path)
            txt_db.save({'date': date, 'model': rec.get('model', ''),
                'station': _normalize_station(rec.get('station', '') or rec.get('st', '')),
                'line': rec.get('line', ''), 'sheetId': rec.get('sheetId', ''),
                'tc': rec.get('tc', ''), 'x': str(rec.get('x', '') or ''),
                'y': str(rec.get('y', '') or ''), 'lv': rec.get('lv', '') or rec.get('level', ''),
                'note': rec.get('note', ''), 'empId': rec.get('empId', ''),
                'photoUrls': photo_urls, 'createdAt': datetime.now().isoformat()})
            saved += 1
        log(f'✅ 成功儲存 {saved} 筆')
        handler.send_json({'success': True, 'count': saved})
    except Exception as e:
        log(f'❌ 儲存失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def check_isnew(handler, data):
    try:
        records   = data.get('records', [])
        tolerance = float(data.get('tolerance', 0.2))
        if not records:
            handler.send_json({'success': True, 'results': []}); return
        prev_needed = set()
        for rec in records:
            st = _normalize_station(rec.get('station', ''))
            if st in STATION_ORDER:
                idx = STATION_ORDER.index(st)
                if idx > 0: prev_needed.add(STATION_ORDER[idx - 1])
        prev_data = {}
        for station in prev_needed:
            prev_data[station] = txt_db.search({'station': station}, '', '')
        results = []
        for rec in records:
            st = _normalize_station(rec.get('station', ''))
            if st not in STATION_ORDER:
                results.append({'isNew': None, 'reason': '非標準站點'}); continue
            idx = STATION_ORDER.index(st)
            if idx == 0:
                results.append({'isNew': True, 'reason': '首站'}); continue
            prev_st   = STATION_ORDER[idx - 1]
            prev_recs = prev_data.get(prev_st, [])
            sid = rec.get('sheetId', '').strip()
            if not sid:
                results.append({'isNew': None, 'reason': '無片號'}); continue
            try:
                y_lo, y_hi = _parse_y_range(rec.get('y', ''))
            except Exception:
                y_lo = y_hi = None
            matched = [p for p in prev_recs if p.get('sheetId', '').strip() == sid]
            if not matched:
                results.append({'isNew': True, 'reason': f'{prev_st} 無此片號'}); continue
            if y_lo is None:
                results.append({'isNew': None, 'reason': '無 Y 值'}); continue
            found = False
            for p in matched:
                try:
                    py_lo, py_hi = _parse_y_range(p.get('y', ''))
                    if py_lo is None: continue
                    if abs(py_lo - y_lo) <= tolerance or abs(py_hi - y_hi) <= tolerance:
                        found = True; break
                except Exception:
                    continue
            results.append({'isNew': not found, 'reason': '比對完成'})
        handler.send_json({'success': True, 'results': results})
    except Exception as e:
        log(f'❌ isnew 比對失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def update_record(handler, data):
    try:
        created_at    = data.get('createdAt', '')
        original_date = data.get('originalDate', '')
        record        = data.get('record', {})
        new_photos    = data.get('newPhotos', [])
        if not created_at or not record:
            handler.send_json({'success': False, 'error': '缺少必要欄位'}, 400); return
        existing = txt_db.find_by_created_at(created_at, original_date)
        if not existing:
            handler.send_json({'success': False, 'error': '找不到原始紀錄'}, 404); return
        photo_urls = list(record.get('photoUrls', []))
        for photo in new_photos:
            if 'dataUrl' not in photo: continue
            raw = decode_base64_image(photo['dataUrl'])
            if raw:
                ok, path = save_photo_from_bytes(raw, photo['name'], PHOTOS_DIR, record.get('date', ''))
                if ok: photo_urls.append(path)
        record['photoUrls'] = photo_urls
        ok, msg = txt_db.update(created_at, original_date, record)
        if ok:
            handler.send_json({'success': True, 'record': record})
        else:
            handler.send_json({'success': False, 'error': msg}, 500)
    except Exception as e:
        log(f'❌ 更新失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def delete_records(handler, data):
    try:
        deletions = data.get('deletions', [])
        if not deletions:
            handler.send_json({'success': False, 'error': '無刪除項目'}, 400); return
        deleted = errors = 0
        for item in deletions:
            ok, msg = txt_db.delete(item.get('createdAt', ''), item.get('date', ''))
            if ok: deleted += 1
            else:  errors += 1
        handler.send_json({'success': True, 'count': deleted, 'errors': errors})
    except Exception as e:
        log(f'❌ 刪除失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def import_records(handler, data):
    try:
        records = data.get('records', [])
        if not records:
            handler.send_json({'success': False, 'error': '無匯入記錄'}, 400); return
        saved = skipped = 0
        for rec in records:
            date = rec.get('date', '').strip()
            if not date or len(date) < 8: skipped += 1; continue
            try:
                txt_db.save({'date': date, 'model': rec.get('model', ''),
                    'station': _normalize_station(rec.get('station', '')),
                    'line': rec.get('line', ''), 'sheetId': rec.get('sheetId', ''),
                    'tc': rec.get('tc', ''), 'x': str(rec.get('x', '') or ''),
                    'y': str(rec.get('y', '') or ''), 'lv': rec.get('lv', '') or rec.get('level', ''),
                    'note': rec.get('note', ''), 'empId': rec.get('empId', ''),
                    'photoUrls': rec.get('photoUrls', []), 'createdAt': datetime.now().isoformat()})
                saved += 1
            except Exception as row_err:
                log(f'⚠️ 單筆匯入失敗：{row_err}'); skipped += 1
        handler.send_json({'success': True, 'count': saved, 'skipped': skipped})
    except Exception as e:
        log(f'❌ 匯入失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def preview_excel_with_pics(handler, data):
    import openpyxl, uuid, base64
    try:
        b64 = data.get('excel_base64', '')
        if not b64:
            handler.send_json({'success': False, 'error': '缺少 excel_base64'}, 400); return
        raw_bytes  = base64.b64decode(b64)
        cache_dir  = os.path.join(os.path.dirname(TXT_DIR), 'preview_cache')
        os.makedirs(cache_dir, exist_ok=True)
        pid        = str(uuid.uuid4())
        cache_path = os.path.join(cache_dir, pid + '.xlsx')
        with open(cache_path, 'wb') as fw: fw.write(raw_bytes)
        wb = openpyxl.load_workbook(cache_path, data_only=True)
        ws = wb.active
        hdrs = [str(cell.value).strip() if cell.value else f'col_{i}'
                for i, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))]
        rows = [{hdrs[i]: (str(v).strip() if v is not None else '')
                 for i, v in enumerate(row) if i < len(hdrs)}
                for row in ws.iter_rows(min_row=2, values_only=True)
                if any(v is not None for v in row)]
        row_images = {}
        # ★ 改用 zip 直讀法
        _zip_imgs = _get_images_from_xlsx_zip(cache_path)
        import base64 as _b64c
        for _rk, _blist in _zip_imgs.items():
            row_images[_rk] = ['data:image/jpeg;base64,' + _b64c.b64encode(_b).decode() for _b in _blist]
        wb.close()
        handler.send_json({'success': True, 'preview_id': pid, 'headers': hdrs,
                           'rows': rows[:20], 'total': len(rows), 'row_images': row_images})
    except Exception as e:
        log(f'❌ FEOL 預覽失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)

def import_excel_with_pics(handler, data):
    import openpyxl, re as _re
    def _fmt(val):
        if not val: return ''
        if hasattr(val, 'strftime'): return val.strftime('%Y-%m-%d')
        s = str(val).strip()
        if 'T' in s: return s[:10]
        if len(s) >= 10 and s[4] == '-' and s[7] == '-' and ' ' in s: return s[:10]
        m = _re.match(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$', s)
        if m: return f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
        return s
    def _get(vals, hdrs, *names):
        for n in names:
            if n in hdrs:
                i = hdrs.index(n)
                if i < len(vals) and vals[i] is not None and str(vals[i]).strip():
                    return str(vals[i]).strip()
        return ''
    def _ext(b):
        if b[:2] == b'\xff\xd8': return 'jpg'
        if b[:4] == b'\x89PNG': return 'png'
        return 'jpg'
    try:
        pid = data.get('preview_id', '')
        if not pid:
            handler.send_json({'success': False, 'error': '缺少 preview_id'}, 400); return
        cache_dir  = os.path.join(os.path.dirname(TXT_DIR), 'preview_cache')
        cache_path = os.path.join(cache_dir, pid + '.xlsx')
        if not os.path.exists(cache_path):
            handler.send_json({'success': False, 'error': '預覽已過期，請重新上傳'}, 404); return
        wb = openpyxl.load_workbook(cache_path, data_only=True)
        ws = wb.active
        hdrs = [str(cell.value).strip() if cell.value else f'col_{i}'
                for i, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))]
        images_by_row = {}
        # ★ 改用 zip 直讀法
        _zip_imgs2 = _get_images_from_xlsx_zip(cache_path)
        for _rk2, _blist2 in _zip_imgs2.items():
            images_by_row[int(_rk2)] = [(10, _b2) for _b2 in _blist2]
        os.makedirs(PHOTOS_DIR, exist_ok=True)
        saved = skipped = photo_saved = 0
        for data_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
            vals = list(row)
            if all(v is None for v in vals): continue
            date    = _fmt(_get(vals, hdrs, '日期', 'date', 'Date', 'MFG_DAY'))
            if not date or len(date) < 8: skipped += 1; continue
            model   = _get(vals, hdrs, 'Model', 'model', 'MODEL', '機型', 'PRODUCT_CODE', 'product_code')
            station = _normalize_station(_get(vals, hdrs, '站點', 'station', 'Station'))
            line    = _get(vals, hdrs, 'LINE', 'Line', 'line', 'LINE_ID', 'line_id')
            sheetId = _get(vals, hdrs, 'Sheet ID', 'SheetID', 'sheetId', 'sheet_id', '片號', 'SHEET_ID')
            tc      = _get(vals, hdrs, 'T/C側', 'T/C 側', 'TC', 'tc', 'T/C', 't/c')
            x       = _get(vals, hdrs, 'X(cm)', 'X (cm)', 'x', 'X')
            y       = _get(vals, hdrs, 'Y(cm)', 'Y (cm)', 'y', 'Y')
            lv      = _get(vals, hdrs, '程度(LV)', '程度 (LV)', '程度', 'LV', 'lv', 'level')
            note    = _get(vals, hdrs, 'Note', 'note', '備註', 'NOTE')
            empId   = _get(vals, hdrs, '工號', 'empId', 'EMP_ID', 'emp_id', 'Emp_ID')
            dt  = date.replace('-', ''); st = station.replace(' ', '')
            pos = f"X{x}_Y{y}" if (x or y) else 'pos'
            photo_urls = []
            for seq, (col_idx, img_bytes) in enumerate(images_by_row.get(data_idx, []), start=1):
                try:
                    ext_  = _ext(img_bytes)
                    fname = f"{dt}_{model}_{st}_{tc}_{pos}_{seq}.{ext_}"
                    fpath = os.path.join(PHOTOS_DIR, fname)
                    with open(fpath, 'wb') as fw: fw.write(img_bytes)
                    photo_urls.append(fpath); photo_saved += 1
                except: pass
            txt_db.save({'date': date, 'model': model, 'station': station, 'line': line,
                         'sheetId': sheetId, 'tc': tc, 'x': x, 'y': y, 'lv': lv, 'note': note,
                         'empId': empId, 'photoUrls': photo_urls, 'createdAt': datetime.now().isoformat()})
            saved += 1
        wb.close()
        try: os.remove(cache_path)
        except: pass
        handler.send_json({'success': True, 'count': saved, 'photos': photo_saved, 'skipped': skipped})
    except Exception as e:
        log(f'❌ 匯入失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def _get_img_bytes(img):
    """嘗試多種方式取得圖片 bytes（相容各版 openpyxl）"""
    try:
        d = getattr(img, '_data', None)
        b = d() if callable(d) else d
        if isinstance(b, bytes) and b: return b
    except: pass
    try:
        ref = getattr(img, 'ref', None)
        if ref is None: return None
        if hasattr(ref, 'getvalue'):
            b = ref.getvalue(); return b if b else None
        if hasattr(ref, 'read'):
            try: ref.seek(0)
            except: pass
            b = ref.read(); return b if b else None
        if isinstance(ref, bytes) and ref: return ref
    except: pass
    return None


def _get_images_from_xlsx_zip(xlsx_path):
    """
    直接解析 xlsx zip 取圖片 bytes 及所在 row（0-indexed）。
    不依賴 ws._images，相容所有 openpyxl 版本。
    回傳 {str(row): [bytes, ...]}
    """
    import zipfile, xml.etree.ElementTree as ET
    XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    A   = "http://schemas.openxmlformats.org/drawingml/2006/main"
    R   = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    result = {}
    try:
        with zipfile.ZipFile(xlsx_path, "r") as zf:
            names = set(zf.namelist())
            draw_files = sorted([n for n in names
                                 if n.startswith("xl/drawings/drawing") and n.endswith(".xml")])
            log(f"[圖片] zip drawing_files={draw_files}")
            for dxml in draw_files:
                drels = dxml.replace("drawings/", "drawings/_rels/") + ".rels"
                rid_map = {}
                if drels in names:
                    for rel in ET.fromstring(zf.read(drels)):
                        rid    = rel.get("Id", "")
                        target = rel.get("Target", "")
                        if "media" in target:
                            rid_map[rid] = "xl/media/" + target.split("/")[-1]
                log(f"[圖片] zip rid_map={rid_map}")
                if not rid_map:
                    continue
                draw_root = ET.fromstring(zf.read(dxml))
                for anchor_elem in draw_root:
                    from_tag = anchor_elem.find(f"{{{XDR}}}from")
                    if from_tag is None:
                        continue
                    row_tag = from_tag.find(f"{{{XDR}}}row")
                    if row_tag is None:
                        continue
                    try:
                        row_val = int(row_tag.text)
                    except Exception:
                        continue
                    rid = None
                    for blip in anchor_elem.iter(f"{{{A}}}blip"):
                        rid = blip.get(f"{{{R}}}embed")
                        break
                    if not rid or rid not in rid_map:
                        continue
                    zip_path = rid_map[rid]
                    if zip_path not in names:
                        continue
                    img_bytes = zf.read(zip_path)
                    log(f"[圖片] zip row={row_val} rid={rid} size={len(img_bytes)}")
                    result.setdefault(str(row_val), []).append(img_bytes)
    except Exception as e:
        log(f"[圖片] zip 解析失敗: {e}")
    return result


def bulk_update(handler, data):
    """統一修改：將搜尋結果中所有記錄的某欄位改為指定值"""
    try:
        records = data.get('records', [])
        field   = data.get('field', '')
        value   = data.get('value', '')
        ALLOWED = {'line', 'model', 'station', 'tc', 'lv', 'note', 'empId'}
        if field not in ALLOWED:
            handler.send_json({'success': False, 'error': f'不允許修改欄位：{field}'}, 400); return
        if not records:
            handler.send_json({'success': False, 'error': '無記錄'}, 400); return
        updated = errors = 0
        for rec in records:
            created_at = rec.get('createdAt', '')
            date       = rec.get('date', '')
            if not created_at or not date: errors += 1; continue
            new_rec = dict(rec)
            new_rec[field]       = value
            new_rec['updatedAt'] = datetime.now().isoformat()
            ok, msg = txt_db.update(created_at, date, new_rec)
            if ok:  updated += 1
            else:   errors  += 1; log(f'⚠️ 統一修改單筆失敗：{msg}')
        log(f'✅ 統一修改: field={field}, value={value}, updated={updated}, errors={errors}')
        handler.send_json({'success': True, 'count': updated, 'errors': errors})
    except Exception as e:
        log(f'❌ 統一修改失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def machine_map_image(handler):
    """服務底圖 M27.jpg"""
    map_path = '\\\\tw100049089\\00_MyAgent\\Roller Mark 背污逐站檢\\M27.jpg'
    try:
        if not os.path.exists(map_path):
            handler.send_json({'error': '底圖不存在', 'path': map_path}, 404); return
        with open(map_path, 'rb') as f:
            data = f.read()
        handler.send_response(200)
        handler.send_header('Content-Type', 'image/jpeg')
        handler.send_header('Content-Length', str(len(data)))
        handler.send_header('Access-Control-Allow-Origin', '*')
        handler.send_header('Cache-Control', 'public, max-age=86400')
        handler.end_headers()
        handler.wfile.write(data)
        log('✅ M27.jpg 底圖已服務')
    except Exception as e:
        log(f'❌ 底圖讀取失敗：{e}')
        handler.send_error(500)


def machine_all_positions(handler, params):
    """查詢同 tool_id + tool_name 的全部位置"""
    try:
        tool_id   = params.get('tool_id', '')
        tool_name = params.get('tool_name', '')
        if not machine_db.is_loaded():
            handler.send_json({'success': False, 'error': '機台資料未載入'}, 503); return
        records = [r for r in machine_db.records
                   if r.get('tool_id', '') == tool_id
                   and r.get('tool_name', '') == tool_name]
        log(f'機台位置查詢: tool_id={tool_id} tool_name={tool_name} 共{len(records)}筆')
        handler.send_json({'success': True, 'records': records})
    except Exception as e:
        log(f'❌ 機台位置查詢失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)
