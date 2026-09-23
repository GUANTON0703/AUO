#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
beol_handler 202507251300.py — BEOL API 處理器
修改說明：
  - 202507251115：wb.close() 修 WinError 32
  - 202507251200：meta.json 改存路徑，修 bytes JSON serializable
  - 202507251300：三個 bug 一次修
      Bug1 照片預覽顯示「—」：row_images[key] 改為平坦 URL 陣列
           （前端 _showPicPreview 期待 imgs.map(url=>...) 而非 {photo,om} 物件）
      Bug2 匯入後搜尋不到：日期 "2026-08-17 00:00:00" 截取前10碼存 "2026-08-17"
      Bug3 X/Y 欄位對應失敗：新增 'X (cm)'/'Y (cm)' 到 g() key 清單
"""

import os
from datetime import datetime

from config_beol import BEOL_TXT_DIR, BEOL_PHOTOS_DIR
from utils.txt_helper   import TxtHelper
from utils.network_path import decode_base64_image, save_photo_from_bytes

beol_txt_db = TxtHelper(BEOL_TXT_DIR)


def log(msg, level='INFO'):
    print(f'[{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}] [BEOL/{level}] {msg}')


def _fmt_date(s):
    """把各種日期格式統一成 YYYY-MM-DD，處理 openpyxl 回傳的 datetime 字串"""
    if not s:
        return ''
    s = str(s).strip()
    # datetime 物件 str → "2026-08-17 00:00:00" 或 "2026-08-17T00:00:00"
    if len(s) >= 10 and (s[10] == ' ' or s[10] == 'T'):
        return s[:10]
    return s


def _beol_extract_images_from_xlsx(xlsx_path):
    import zipfile, xml.etree.ElementTree as ET

    PHOTO_COL = 11
    OM_COL    = 13
    images_by_row = {}

    try:
        with zipfile.ZipFile(xlsx_path, 'r') as zf:
            names = zf.namelist()
            media = {}
            for n in names:
                if n.startswith('xl/media/'):
                    fname = n.split('/')[-1]
                    media[fname] = zf.read(n)
            log(f'zipfile：找到 {len(media)} 個媒體檔：{list(media.keys())}')
            if not media:
                return images_by_row
            drawing_files = [n for n in names if n.startswith('xl/drawings/drawing') and n.endswith('.xml')]
            log(f'zipfile：找到 drawing 檔：{drawing_files}')
            for drw_path in drawing_files:
                drw_name  = drw_path.split('/')[-1]
                rels_path = f'xl/drawings/_rels/{drw_name}.rels'
                if rels_path not in names:
                    log(f'  找不到 rels：{rels_path}')
                    continue
                rels_xml  = zf.read(rels_path)
                rels_root = ET.fromstring(rels_xml)
                rid_to_media = {}
                for rel in rels_root:
                    rid   = rel.get('Id', '')
                    tgt   = rel.get('Target', '')
                    fname = tgt.split('/')[-1]
                    if fname in media:
                        rid_to_media[rid] = fname
                log(f'  rels 對應：{rid_to_media}')
                drw_xml  = zf.read(drw_path)
                drw_root = ET.fromstring(drw_xml)
                NS = {
                    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
                    'a':   'http://schemas.openxmlformats.org/drawingml/2006/main',
                    'r':   'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
                }
                for anchor in drw_root:
                    tag = anchor.tag.split('}')[-1]
                    if tag not in ('twoCellAnchor', 'oneCellAnchor'):
                        continue
                    from_el = anchor.find('xdr:from', NS)
                    if from_el is None:
                        continue
                    row_el = from_el.find('xdr:row', NS)
                    col_el = from_el.find('xdr:col', NS)
                    if row_el is None or col_el is None:
                        continue
                    try:
                        row = int(row_el.text)
                        col = int(col_el.text)
                    except (ValueError, TypeError):
                        continue
                    pic  = anchor.find('.//xdr:pic', NS)
                    if pic is None:
                        continue
                    blip = pic.find('.//a:blip', NS)
                    if blip is None:
                        continue
                    rid = blip.get(
                        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', '')
                    if not rid or rid not in rid_to_media:
                        continue
                    img_bytes = media[rid_to_media[rid]]
                    log(f'  圖片：row={row} col={col} rId={rid} size={len(img_bytes)}')
                    if col < 10:
                        continue
                    is_om = abs(col - OM_COL) < abs(col - PHOTO_COL)
                    if row not in images_by_row:
                        images_by_row[row] = {'photo': [], 'om': []}
                    if is_om:
                        images_by_row[row]['om'].append(img_bytes)
                    else:
                        images_by_row[row]['photo'].append(img_bytes)
    except Exception as e:
        log(f'❌ zipfile 解析圖片失敗：{e}')

    log(f'zipfile 解析完成：{len(images_by_row)} 列有圖片，rows={list(images_by_row.keys())}')
    return images_by_row


def beol_save(handler, data):
    try:
        date    = data.get('date', '')
        records = data.get('records', [])
        if not date or not records:
            handler.send_json({'success': False, 'error': '日期或紀錄缺失'}, 400); return
        saved = 0
        for rec in records:
            photo_urls = []
            for photo in rec.get('photos', []):
                if 'dataUrl' not in photo: continue
                raw = decode_base64_image(photo['dataUrl'])
                if raw:
                    ok, path = save_photo_from_bytes(raw, photo['name'], BEOL_PHOTOS_DIR, date)
                    if ok: photo_urls.append(path)
            om_photo_urls = []
            for photo in rec.get('omPhotos', []):
                if 'dataUrl' not in photo: continue
                raw = decode_base64_image(photo['dataUrl'])
                if raw:
                    ok, path = save_photo_from_bytes(raw, photo['name'], BEOL_PHOTOS_DIR, date)
                    if ok: om_photo_urls.append(path)
            beol_txt_db.save({
                'date': date, 'model': data.get('model', ''),
                'station': data.get('station', ''), 'line': data.get('line', ''),
                'empId': data.get('empId', ''),
                'midId': rec.get('midId', ''),
                'chipId': rec.get('chipId', ''), 'tc': rec.get('tc', ''),
                'x': str(rec.get('x', '') or ''), 'y': str(rec.get('y', '') or ''),
                'morphology': rec.get('morphology', ''), 'level': rec.get('level', ''),
                'ito': rec.get('ito', ''), 'note': rec.get('note', ''),
                'photoUrls': photo_urls, 'omPhotoUrls': om_photo_urls,
                'createdAt': datetime.now().isoformat()
            })
            saved += 1
        log(f'✅ 儲存 {saved} 筆')
        handler.send_json({'success': True, 'count': saved})
    except Exception as e:
        log(f'❌ 儲存失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def beol_search(handler, params):
    try:
        filters = {k: params[k] for k in
                   ['model', 'station', 'line', 'tc', 'level', 'empId',
                    'midId', 'chipId', 'morphology', 'ito']
                   if params.get(k)}
        records = beol_txt_db.search(filters, params.get('dateFrom', ''), params.get('dateTo', ''))
        records.sort(key=lambda r: (r.get('date', ''), r.get('createdAt', '')), reverse=True)
        handler.send_json({'success': True, 'records': records})
    except Exception as e:
        log(f'❌ 搜尋失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def beol_download_template(handler):
    template_path = r'\\tw100049089\ML8BC1\逐站檢資料\BEOL_SAMPLE_01.xlsm'
    try:
        if not os.path.exists(template_path):
            handler.send_json({'success': False, 'error': '範例檔案不存在'}, 404); return
        with open(template_path, 'rb') as f:
            data = f.read()
        handler.send_response(200)
        handler.send_header('Content-Type', 'application/vnd.ms-excel.sheet.macroEnabled.12')
        handler.send_header('Content-Disposition', 'attachment; filename="BEOL_SAMPLE_01.xlsm"')
        handler.send_header('Content-Length', str(len(data)))
        handler.send_header('Access-Control-Allow-Origin', '*')
        handler.end_headers()
        try:
            handler.wfile.write(data)
            log('✅ BEOL 範例 Excel 已下載')
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError, OSError):
            log('⚠️ BEOL 範例下載連線中斷')
    except Exception as e:
        log(f'❌ BEOL 範例下載失敗：{e}')


def beol_export(handler, params):
    try:
        filters = {k: params[k] for k in
                   ['model', 'station', 'line', 'tc', 'level', 'empId'] if params.get(k)}
        records = beol_txt_db.search(filters, params.get('dateFrom', ''), params.get('dateTo', ''))
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = 'BEOL背汙檢'
        ws.append(['日期', 'Model', '站點', 'LINE', '中片ID', 'CHIP ID', 'T/C',
                   'X(mm)', 'Y(mm)', '成像敘述', '程度', 'ITO', 'Note', '工號'])
        for r in records:
            ws.append([r.get('date',''), r.get('model',''), r.get('station',''), r.get('line',''),
                       r.get('midId',''), r.get('chipId',''), r.get('tc',''),
                       r.get('x',''), r.get('y',''), r.get('morphology',''),
                       r.get('level',''), r.get('ito',''), r.get('note',''), r.get('empId','')])
        from io import BytesIO
        buf = BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        handler.send_response(200)
        handler.send_header('Content-Type',
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        handler.send_header('Content-Disposition', 'attachment; filename="BEOL_export.xlsx"')
        handler.send_header('Content-Length', str(len(data)))
        handler.send_header('Access-Control-Allow-Origin', '*')
        handler.end_headers()
        try:
            handler.wfile.write(data)
            log('✅ BEOL 匯出完成')
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError, OSError):
            log('⚠️ BEOL 匯出連線中斷')
    except Exception as e:
        log(f'❌ BEOL 匯出失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def beol_delete(handler, data):
    try:
        items = data.get('deletions', [])
        if not items:
            handler.send_json({'success': False, 'error': '無項目'}, 400); return
        deleted = 0
        for item in items:
            ok = beol_txt_db.delete(item.get('createdAt',''), item.get('date',''))
            if ok: deleted += 1
        log(f'✅ 刪除 {deleted} 筆')
        handler.send_json({'success': True, 'deleted': deleted})
    except Exception as e:
        log(f'❌ 刪除失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def beol_preview_excel_with_pics(handler, data):
    try:
        import base64, tempfile, json, uuid
        b64 = data.get('excel_base64', '')
        if not b64:
            handler.send_json({'success': False, 'error': '無 Excel 資料'}, 400); return
        xlsx_bytes = base64.b64decode(b64)
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp.write(xlsx_bytes)
            tmp_path = tmp.name
        try:
            from openpyxl import load_workbook
            wb = load_workbook(tmp_path, read_only=True, data_only=True)
            ws = wb.active
            headers = []
            rows_data = []
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                if ri == 0:
                    headers = [str(c) if c is not None else '' for c in row]
                else:
                    rows_data.append({
                        headers[ci]: (str(v) if v is not None else '')
                        for ci, v in enumerate(row)
                    })
            wb.close()  # 釋放 Windows 檔案鎖

            images_by_row = _beol_extract_images_from_xlsx(tmp_path)

            preview_id  = str(uuid.uuid4())
            preview_dir = os.path.join(BEOL_PHOTOS_DIR, 'previews', preview_id)
            os.makedirs(preview_dir, exist_ok=True)

            host = handler.headers.get('Host', '127.0.0.1:8888')

            # ★ Bug1 修正：row_images[key] 改為平坦 URL 陣列（photo + om 合併）
            #   前端 _showPicPreview 期待 imgs.map(url=>...) 而非 {photo,om} 物件
            # ★ img_paths[key] 仍保留 {photo,om} 分開，供 import 時區分用
            row_images = {}   # {key: [url, url, ...]}  ← 平坦陣列給前端
            img_paths  = {}   # {key: {photo:[path], om:[path]}}  ← 路徑給 meta.json

            for row_idx, imgs in images_by_row.items():
                key = str(row_idx)
                row_images[key] = []
                img_paths[key]  = {'photo': [], 'om': []}

                for i, img_bytes in enumerate(imgs.get('photo', [])):
                    fname = f'photo_{row_idx}_{i}.png'
                    fpath = os.path.join(preview_dir, fname)
                    with open(fpath, 'wb') as ff:
                        ff.write(img_bytes)
                    url = f'http://{host}/api/photo?path={fpath}'
                    row_images[key].append(url)          # ← 平坦加入
                    img_paths[key]['photo'].append(fpath)

                for i, img_bytes in enumerate(imgs.get('om', [])):
                    fname = f'om_{row_idx}_{i}.png'
                    fpath = os.path.join(preview_dir, fname)
                    with open(fpath, 'wb') as ff:
                        ff.write(img_bytes)
                    url = f'http://{host}/api/photo?path={fpath}'
                    row_images[key].append(url)          # ← 平坦加入
                    img_paths[key]['om'].append(fpath)

            meta_path = os.path.join(preview_dir, 'meta.json')
            with open(meta_path, 'w', encoding='utf-8') as ff:
                json.dump({
                    'headers':     headers,
                    'rows':        rows_data,
                    'img_paths':   img_paths,
                    'preview_dir': preview_dir
                }, ff, ensure_ascii=False)

            handler.send_json({'success': True, 'preview_id': preview_id,
                               'headers': headers, 'rows': rows_data,
                               'row_images': row_images, 'total': len(rows_data)})
        finally:
            os.unlink(tmp_path)
    except Exception as e:
        log(f'❌ BEOL preview_excel_with_pics 失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)


def beol_import_excel_with_pics(handler, data):
    try:
        import json
        preview_id = data.get('preview_id', '')
        if not preview_id:
            handler.send_json({'success': False, 'error': '無 preview_id'}, 400); return
        preview_dir = os.path.join(BEOL_PHOTOS_DIR, 'previews', preview_id)
        meta_path = os.path.join(preview_dir, 'meta.json')
        if not os.path.exists(meta_path):
            handler.send_json({'success': False, 'error': 'Preview 不存在或已過期'}, 404); return
        with open(meta_path, 'r', encoding='utf-8') as ff:
            meta = json.load(ff)
        headers   = meta['headers']
        rows      = meta['rows']
        img_paths = meta.get('img_paths', {})

        saved = 0; photos_saved = 0; skipped = 0

        # ★ Bug3 修正：X/Y 欄位對應新增 'X (cm)'/'Y (cm)'
        g = lambda row, keys: next((str(row[k]) for k in keys if row.get(k) not in (None,'') ), '')

        for ri, row in enumerate(rows):
            # ★ Bug2 修正：日期截取前10碼，避免 "2026-08-17 00:00:00" 比對失敗
            date = _fmt_date(g(row, ['日期', 'date', 'Date', 'DATE']))
            chip = g(row, ['CHIP ID', 'chipId', 'chip_id'])
            if not date and not chip:
                skipped += 1; continue

            photo_urls = []; om_photo_urls = []
            # img_paths key 是 drawing row index（0-based），
            # 第 ri 筆資料（ri=0起）對應 drawing row = ri+1（header 佔 row 0）
            row_paths = img_paths.get(str(ri + 1), {})

            for fpath in row_paths.get('photo', []):
                if not os.path.exists(fpath):
                    continue
                with open(fpath, 'rb') as ff:
                    img_bytes = ff.read()
                ok, path = save_photo_from_bytes(
                    img_bytes, f'import_{ri+1}_ph.png', BEOL_PHOTOS_DIR, date)
                if ok: photo_urls.append(path); photos_saved += 1

            for fpath in row_paths.get('om', []):
                if not os.path.exists(fpath):
                    continue
                with open(fpath, 'rb') as ff:
                    img_bytes = ff.read()
                ok, path = save_photo_from_bytes(
                    img_bytes, f'import_{ri+1}_om.png', BEOL_PHOTOS_DIR, date)
                if ok: om_photo_urls.append(path); photos_saved += 1

            beol_txt_db.save({
                'date':      date,
                'model':     g(row, ['Model', 'model', 'MODEL']),
                'station':   g(row, ['站點', 'station', 'Station']),
                'line':      g(row, ['LINE', 'Line', 'line']),
                'midId':     g(row, ['中片ID', 'midId']),
                'chipId':    chip,
                'tc':        g(row, ['T/C', 'T/C側', 'TC', 'tc']),
                # ★ Bug3 修正：加上 'X (cm)'/'Y (cm)' 對應
                'x':         g(row, ['X(mm)', 'X (mm)', 'X (cm)', 'X(cm)', 'x', 'X']),
                'y':         g(row, ['Y(mm)', 'Y (mm)', 'Y (cm)', 'Y(cm)', 'y', 'Y']),
                'morphology': g(row, ['成像敘述', 'morphology']),
                'level':     g(row, ['程度', 'level', 'Level']),
                'ito':       g(row, ['ITO', 'ito']),
                'note':      g(row, ['Note', 'note', 'NOTE', '備註']),
                'empId':     g(row, ['工號', 'empId']),
                'photoUrls':   photo_urls,
                'omPhotoUrls': om_photo_urls,
                'createdAt': datetime.now().isoformat()
            })
            saved += 1

        log(f'✅ BEOL Excel 匯入 {saved} 筆，圖片 {photos_saved} 張，跳過 {skipped} 筆')
        handler.send_json({'success': True, 'count': saved,
                           'photos': photos_saved, 'skipped': skipped})
    except Exception as e:
        log(f'❌ BEOL import_excel_with_pics 失敗：{e}')
        handler.send_json({'success': False, 'error': str(e)}, 500)
